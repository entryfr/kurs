from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

from docx import Document
from jinja2 import Environment, FileSystemLoader

logger = logging.getLogger(__name__)


def _generate_pdf_act(data: dict[str, Any], pdf_output_path: Path) -> Path | None:
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
    except ImportError:
        logger.warning(
            "reportlab не установлен, PDF-версия акта не будет создана: %s",
            pdf_output_path,
        )
        return None

    pdf_output_path.parent.mkdir(parents=True, exist_ok=True)
    c = canvas.Canvas(str(pdf_output_path), pagesize=A4)
    _, height = A4
    y = height - 40

    lines = [
        "АКТ ВЫЯВЛЕННЫХ НЕУЧТЁННЫХ КОММУНИКАЦИЙ",
        f"Дата: {data.get('date', '-')}",
        f"Профиль норм: {data.get('profile', '-')}",
        f"Количество аномалий: {data.get('anomalies', 0)}",
        "",
        "Краткая сводка по аномалиям:",
    ]
    for row in data.get("anomaly_rows", [])[:25]:
        lines.append(
            f"#{row.get('anomaly_id', '-')}: {row.get('risk_class', '-')}, "
            f"P={row.get('risk_probability', '-')}, "
            f"тип={row.get('utility_type', '-')}, "
            f"рекомендация={row.get('recommendation', '-')}"
        )

    for line in lines:
        if y < 40:
            c.showPage()
            y = height - 40
        c.drawString(40, y, line[:150])
        y -= 16

    c.save()
    return pdf_output_path


def generate_act(
    data: dict[str, Any],
    template_path: str = "templates/act_template.jinja2",
    output_path: str = "act.docx",
    pdf_output_path: str | None = None,
) -> dict[str, str | None]:
    env = Environment(loader=FileSystemLoader("."))
    template = env.get_template(template_path)
    rendered = template.render(data)

    output = Path(output_path)
    output.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    doc.add_heading("Акт выявленных неучтённых коммуникаций", level=1)
    doc.add_paragraph(rendered)

    meta = doc.add_paragraph()
    meta.add_run(f"Дата: {data.get('date', '-')}\n")
    meta.add_run(f"Профиль норм: {data.get('profile', '-')}\n")
    meta.add_run(f"Количество аномалий: {data.get('anomalies', 0)}\n")

    anomaly_rows = data.get("anomaly_rows", [])
    if anomaly_rows:
        doc.add_heading("Таблица выявленных аномалий", level=2)
        table = doc.add_table(rows=1, cols=8)
        headers = [
            "ID",
            "Риск",
            "Вероятность",
            "Тип",
            "Глубина, м",
            "Длина, м",
            "Координаты",
            "Рекомендация",
        ]
        for idx, header in enumerate(headers):
            table.rows[0].cells[idx].text = header

        for row in anomaly_rows:
            cells = table.add_row().cells
            cells[0].text = str(row.get("anomaly_id", "-"))
            cells[1].text = str(row.get("risk_class", "-"))
            cells[2].text = str(row.get("risk_probability", "-"))
            cells[3].text = str(row.get("utility_type", "-"))
            cells[4].text = str(row.get("depth_m", "-"))
            cells[5].text = str(row.get("length_m", "-"))
            cells[6].text = str(row.get("coordinates", "-"))
            cells[7].text = str(row.get("recommendation", "-"))

    issues = data.get("issues", [])
    if issues:
        doc.add_heading("Замечания и нарушения", level=2)
        for issue in issues:
            doc.add_paragraph(str(issue), style="List Bullet")

    doc.save(str(output))

    pdf_path: Path | None = None
    if pdf_output_path:
        pdf_path = _generate_pdf_act(data, Path(pdf_output_path))

    return {
        "docx_path": str(output),
        "pdf_path": str(pdf_path) if pdf_path is not None else None,
    }